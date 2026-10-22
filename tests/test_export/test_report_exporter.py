"""Comprehensive unit and integration tests for ReportExportEngine and export API endpoints."""

import io
import json

import docx
import pytest
from fastapi.testclient import TestClient

from src.api.job_runner import job_service
from src.api.main import app
from src.core.telemetry import RunCostSummary
from src.export.engine import ReportExportEngine
from src.schemas.api import ResearchStatusResponse
from src.schemas.report import Citation, ReportSection, ResearchReport
from src.validators.rag_scorer import RAGQualityScore


@pytest.fixture
def sample_status_response() -> ResearchStatusResponse:
    """Construct a full-featured mock ResearchStatusResponse with report and telemetry."""
    report = ResearchReport(
        title="Comprehensive Analysis: Multi-Agent Distributed Architectures",
        executive_summary="Multi-agent systems demonstrate superior fault tolerance and specialized execution compared to monolithic workflows.",
        sections=[
            ReportSection(
                title="State Synchronization & Persistence",
                content="Distributed memory utilizes Redis for durable superstep state and pub/sub message broadcasting [cite_1].",
                citation_ids=["cite_1"],
            ),
            ReportSection(
                title="Browser Resilience & Anti-Bot Bypass",
                content="Automated fallback to headless stealth browsers enables reliable data extraction on protected domains [cite_2].",
                citation_ids=["cite_2"],
            ),
        ],
        citations=[
            Citation(
                citation_id="cite_1",
                source_url="https://redis.io/docs/latest/develop/data-types/streams/",
                anchor_url="https://redis.io/docs/latest/develop/data-types/streams/#:~:text=streams",
                verified_claim="Redis streams provide durable append-only log primitives.",
                verbatim_quote="Redis Streams is a data structure that acts like an append-only log.",
                http_status=200,
                is_deep_link=True,
            ),
            Citation(
                citation_id="cite_2",
                source_url="https://playwright.dev/python/docs/intro",
                anchor_url="https://playwright.dev/python/docs/intro#:~:text=Chromium",
                verified_claim="Headless Chromium automates modern dynamic single page applications.",
                verbatim_quote="Playwright provides a single API to automate Chromium, Firefox and WebKit.",
                http_status=200,
                is_deep_link=True,
            ),
        ],
    )

    return ResearchStatusResponse(
        run_id="test-export-run-12345678",
        query="Explain multi-agent state persistence and web scraping resilience",
        status="completed",
        progress_pct=100,
        current_node="writer",
        step_count=12,
        total_tokens=14250,
        findings_count=8,
        report=report,
        rag_score=RAGQualityScore(
            groundedness_score=0.96,
            citation_precision=0.92,
            coverage_score=1.0,
            overall_score=0.95,
            grade="A",
            total_citations=2,
            verified_citations_count=2,
            hallucinated_count=0,
        ),
        cost_summary=RunCostSummary(
            total_cost_usd=0.0142,
            total_input_tokens=11000,
            total_output_tokens=3250,
        ),
        created_at="2026-10-21T06:00:00Z",
    )


def test_export_markdown(sample_status_response: ResearchStatusResponse):
    """Test markdown generation includes title, executive summary, sections, references, and telemetry."""
    md = ReportExportEngine.export_markdown(sample_status_response)

    assert "# Comprehensive Analysis: Multi-Agent Distributed Architectures" in md
    assert "## Executive Summary" in md
    assert "Multi-agent systems demonstrate superior fault tolerance" in md
    assert "## State Synchronization & Persistence" in md
    assert "## References & Verified Citations" in md
    assert "cite_1" in md
    assert "https://redis.io" in md
    assert "Quality & Telemetry Appendix" in md
    assert "Grade A" in md
    assert "$0.01420 USD" in md


def test_export_html(sample_status_response: ResearchStatusResponse):
    """Test HTML generation produces valid markup with styled CSS and proper sections."""
    html_doc = ReportExportEngine.export_html(sample_status_response)

    assert "<!DOCTYPE html>" in html_doc
    assert "Comprehensive Analysis: Multi-Agent Distributed Architectures" in html_doc
    assert "Executive Summary" in html_doc
    assert "State Synchronization &amp; Persistence" in html_doc
    assert "citations-table" in html_doc
    assert (
        "Audit &amp; Verification Ledger" in html_doc or "Audit & Verification Ledger" in html_doc
    )
    assert "HTTP 200" in html_doc


def test_export_docx(sample_status_response: ResearchStatusResponse):
    """Test DOCX export generates a readable Word document with tables and paragraphs."""
    docx_bytes = ReportExportEngine.export_docx(sample_status_response)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 5000

    doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    assert "Comprehensive Analysis: Multi-Agent Distributed Architectures" in full_text
    assert "Executive Summary" in full_text
    assert "State Synchronization & Persistence" in full_text

    # Verify Citations table
    assert len(doc.tables) >= 1
    table_text = " ".join([cell.text for row in doc.tables[0].rows for cell in row.cells])
    assert "cite_1" in table_text
    assert "Redis" in table_text


def test_export_json(sample_status_response: ResearchStatusResponse):
    """Test JSON export returns valid serialized research status payload."""
    json_bytes = ReportExportEngine.export_json(sample_status_response)
    assert isinstance(json_bytes, bytes)

    parsed = json.loads(json_bytes.decode("utf-8"))
    assert parsed["run_id"] == "test-export-run-12345678"
    assert parsed["status"] == "completed"
    assert (
        parsed["report"]["title"] == "Comprehensive Analysis: Multi-Agent Distributed Architectures"
    )


@pytest.mark.asyncio
async def test_export_pdf(sample_status_response: ResearchStatusResponse):
    """Test PDF generation produces valid binary PDF via Playwright."""
    try:
        pdf_bytes = await ReportExportEngine.export_pdf(sample_status_response)
        assert isinstance(pdf_bytes, bytes)
        assert len(pdf_bytes) > 1000
        assert pdf_bytes.startswith(b"%PDF-")
    except Exception as exc:
        if "Executable doesn't exist" in str(exc) or "playwright install" in str(exc):
            pytest.skip(f"Playwright chromium browser not installed: {exc}")
        raise


@pytest.mark.asyncio
async def test_export_dispatch(sample_status_response: ResearchStatusResponse):
    """Test unified export dispatch returns correct content, media-type, and filename."""
    # Test Markdown
    content, media, filename = await ReportExportEngine.export(
        sample_status_response, format="markdown"
    )
    assert media == "text/markdown; charset=utf-8"
    assert filename == "research_report_test-exp.md"
    assert b"Executive Summary" in content

    # Test DOCX
    content, media, filename = await ReportExportEngine.export(
        sample_status_response, format="docx"
    )
    assert media == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert filename == "research_report_test-exp.docx"
    assert len(content) > 5000

    # Test HTML
    content, media, filename = await ReportExportEngine.export(
        sample_status_response, format="html"
    )
    assert media == "text/html; charset=utf-8"
    assert filename == "research_report_test-exp.html"
    assert b"<!DOCTYPE html>" in content

    # Test Unsupported format raises ValueError
    with pytest.raises(ValueError, match="Unsupported export format"):
        await ReportExportEngine.export(sample_status_response, format="invalid_fmt")


@pytest.mark.asyncio
async def test_api_export_endpoint(
    sample_status_response: ResearchStatusResponse, monkeypatch: pytest.MonkeyPatch
):
    """Test FastAPI GET /research/{run_id}/export route integration."""
    client = TestClient(app)

    async def mock_get_status(run_id: str):
        if run_id == sample_status_response.run_id:
            return sample_status_response
        if run_id == "run-in-progress":
            return ResearchStatusResponse(
                run_id="run-in-progress",
                query="test query",
                status="researching",
                progress_pct=40,
                report=None,
            )
        return None

    monkeypatch.setattr(job_service, "get_status", mock_get_status)

    # 1. Export as Markdown
    resp_md = client.get(f"/research/{sample_status_response.run_id}/export?format=md")
    assert resp_md.status_code == 200
    assert "text/markdown" in resp_md.headers["content-type"]
    assert (
        'attachment; filename="research_report_test-exp.md"'
        in resp_md.headers["content-disposition"]
    )
    assert "Executive Summary" in resp_md.text

    # 2. Export as DOCX
    resp_docx = client.get(f"/research/{sample_status_response.run_id}/export?format=docx")
    assert resp_docx.status_code == 200
    assert "wordprocessingml.document" in resp_docx.headers["content-type"]
    assert (
        'attachment; filename="research_report_test-exp.docx"'
        in resp_docx.headers["content-disposition"]
    )
    assert len(resp_docx.content) > 5000

    # 3. Export as HTML
    resp_html = client.get(f"/research/{sample_status_response.run_id}/export?format=html")
    assert resp_html.status_code == 200
    assert "text/html" in resp_html.headers["content-type"]
    assert "<!DOCTYPE html>" in resp_html.text

    # 4. Export as JSON
    resp_json = client.get(f"/research/{sample_status_response.run_id}/export?format=json")
    assert resp_json.status_code == 200
    assert "application/json" in resp_json.headers["content-type"]
    data = resp_json.json()
    assert data["run_id"] == sample_status_response.run_id

    # 5. Invalid format returns 400
    resp_invalid = client.get(
        f"/research/{sample_status_response.run_id}/export?format=unsupported"
    )
    assert resp_invalid.status_code == 400

    # 6. Nonexistent run returns 404
    resp_404 = client.get("/research/nonexistent-run-id/export?format=md")
    assert resp_404.status_code == 404

    # 7. In-progress run with no report returns 400
    resp_progress = client.get("/research/run-in-progress/export?format=pdf")
    assert resp_progress.status_code == 400
    assert "does not have a completed report yet" in resp_progress.json()["error"]
